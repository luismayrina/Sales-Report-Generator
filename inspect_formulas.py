import openpyxl
wb = openpyxl.load_workbook("Sample Reports.xlsx")
ws = wb["Offtake Report Summary"]
for row in [22, 23, 24, 25]:
    name = ws.cell(row=row, column=1).value
    for col in range(2, 12):
        cell = ws.cell(row=row, column=col)
        print(f"{name} Col {col}: {cell.value}")
