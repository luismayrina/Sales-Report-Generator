import openpyxl
wb = openpyxl.load_workbook("Full_Sales_Report.xlsx", data_only=True)
ws = wb["Offtake Report Summary"]

# Common Room is row 2
# January block starts at col 2
actual_col = 3
target_col = 4
py_col = 5
vs_py_col = 6

print(f"Actual: {ws.cell(row=2, column=actual_col).value}")
print(f"Target: {ws.cell(row=2, column=target_col).value}")
print(f"PY: {ws.cell(row=2, column=py_col).value}")
print(f"vs PY: {ws.cell(row=2, column=vs_py_col).value}")
