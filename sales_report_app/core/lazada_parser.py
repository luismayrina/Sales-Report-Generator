from openpyxl import load_workbook
from collections import defaultdict
from .workbook_utils import fval, parse_lazada_date

def load_lazada(path):
    wb = load_workbook(path)
    ws = wb.active
    headers = [c.value for c in ws[1]]
    hi = {h: i for i, h in enumerate(headers)}

    grouped = defaultdict(list)
    for row in ws.iter_rows(min_row=2, values_only=True):
        onum = row[hi["orderNumber"]]
        grouped[onum].append(row)

    orders = []
    for onum, grp in grouped.items():
        first   = grp[0]
        status  = str(first[hi["status"]]).strip().lower()
        created = parse_lazada_date(str(first[hi["createTime"]]))
        delivered = parse_lazada_date(str(first[hi["deliveredDate"]]))

        total_paid = sum(fval(r[hi["paidPrice"]]) for r in grp)
        total_unit = sum(fval(r[hi["unitPrice"]]) for r in grp)
        total_sdisc = sum(fval(r[hi["sellerDiscountTotal"]]) for r in grp)
        total_ship = sum(fval(r[hi["shippingFee"]]) for r in grp)

        is_cancelled = status == "canceled"
        dep_dt = delivered if status in ("delivered", "confirmed") else None

        items = [(str(r[hi["itemName"]]).strip(),
                  1,
                  fval(r[hi["unitPrice"]])) for r in grp]

        orders.append({
            "source":       "Lazada",
            "date":         created,
            "customer":     str(first[hi["customerName"]]).strip(),
            "order_no":     str(onum),
            "gross":        total_unit,
            "shipping":     total_ship or None,
            "discount":     abs(total_sdisc) if total_sdisc < 0 else None,
            "paymongo_fee": None,
            "payment_fee":  None,
            "net":          total_paid,
            "deposit_date": dep_dt,
            "bank":         "LAZADA" if not is_cancelled else "CANCELLED",
            "ref":          None,
            "status":       "cancelled" if is_cancelled else ("paid" if dep_dt else "pending"),
            "items":        items,
        })
    return orders
