import os
import re
from datetime import datetime
import openpyxl
import calendar
from bs4 import BeautifulSoup

try:
    import pdfplumber
except ImportError:
    pdfplumber = None

def clean_amount(val):
    if not val:
        return 0.0
    v = str(val).replace('₱', '').replace(',', '').strip()
    try:
        return float(v)
    except:
        return 0.0

def create_synthetic_order(source, date_obj, amount, items=None, order_no=None):
    if amount == 0 and not items:
        return None
    return {
        "source": source,
        "date": date_obj,
        "customer": f"{source} Customer",
        "order_no": order_no or f"{source[:2].upper()}-{date_obj.strftime('%Y%m%d')}-{abs(hash(str(amount))) % 10000}",
        "gross": amount,
        "shipping": 0.0,
        "discount": 0.0,
        "net": amount,
        "deposit_date": date_obj,
        "bank": source.upper(),
        "status": "paid",
        "items": items or []
    }

def extract_month_from_filename(filename):
    months = ["january", "february", "march", "april", "may", "june", 
              "july", "august", "september", "october", "november", "december",
              "jan", "feb", "mar", "apr", "aug", "sep", "oct", "nov", "dec"]
    name_lower = filename.lower()
    for m in months:
        if m in name_lower:
            try:
                dt = datetime.strptime(m[:3], "%b")
                return dt.month
            except:
                pass
    return None

def parse_tiktok_excel(filepath, current_year=2026):
    txns = []
    try:
        wb = openpyxl.load_workbook(filepath, data_only=True)
        ws = wb.active
        
        headers = [str(c.value).strip() if c.value else "" for c in ws[1]]
        
        try:
            date_idx = headers.index("Created Time")
            amount_idx = headers.index("Order Amount")
            status_idx = headers.index("Order Status")
            prod_name_idx = headers.index("Product Name")
            qty_idx = headers.index("Quantity")
        except ValueError:
            date_idx = 24
            amount_idx = 22
            status_idx = 1
            prod_name_idx = 7
            qty_idx = 9
            
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or not any(row): continue
            
            status = str(row[status_idx]).strip().lower()
            if status in ["canceled", "unpaid"]: continue
            
            date_str = str(row[date_idx])
            dt = None
            try:
                for fmt in ("%m/%d/%Y %I:%M:%S %p", "%Y-%m-%d %H:%M:%S", "%d/%m/%Y %H:%M:%S"):
                    try:
                        dt = datetime.strptime(date_str, fmt)
                        break
                    except ValueError:
                        pass
            except:
                pass
                
            if not dt: continue
            
            amount = clean_amount(row[amount_idx])
            prod_name = str(row[prod_name_idx]).strip() if row[prod_name_idx] else "Tiktok Item"
            qty = clean_amount(row[qty_idx])
            if qty == 0: qty = 1
            
            item_price = amount / qty if qty > 0 else amount
            
            txns.append({
                "source": "Tiktok",
                "date": dt,
                "customer": "Tiktok Customer",
                "order_no": str(row[0]),
                "gross": amount,
                "shipping": 0.0,
                "discount": 0.0,
                "net": amount,
                "deposit_date": dt,
                "bank": "TIKTOK",
                "status": "paid",
                "items": [(prod_name, int(qty), item_price)]
            })
    except Exception as e:
        print(f"Error parsing Tiktok {filepath}: {e}")
    return txns

def parse_simula_excel(filepath, store_name="Simula PH", current_year=2026):
    txns = []
    try:
        wb = openpyxl.load_workbook(filepath, data_only=True)
        for ws in wb.worksheets:
            header_row_idx = 2
            monthly_items = {} 
            
            for row in ws.iter_rows(min_row=header_row_idx + 1, values_only=True):
                if not row or not any(row): continue
                
                date_val = row[0]
                qty_val = row[1]
                prod_val = row[2]
                amount_val = row[3]
                
                if not date_val or not prod_val: continue
                
                amount = clean_amount(amount_val)
                qty = clean_amount(qty_val)
                if qty == 0: qty = 1
                
                dt = None
                if isinstance(date_val, datetime):
                    dt = date_val
                elif date_val:
                    date_str = str(date_val).strip()
                    if "+" in date_str:
                        date_str = date_str.split("+")[0].strip()
                    for fmt in ("%Y-%m-%d %H:%M:%S", "%d %b %Y %I:%M %p", "%m/%d/%Y", "%d/%m/%Y", "%Y-%m-%d"):
                        try:
                            dt = datetime.strptime(date_str, fmt)
                            break
                        except ValueError:
                            pass
                            
                if dt:
                    month = dt.month
                    if month not in monthly_items:
                        monthly_items[month] = []
                    item_price = amount / qty if qty > 0 else amount
                    monthly_items[month].append((str(prod_val).strip(), int(qty), item_price))
                    
            for m, items in monthly_items.items():
                last_day = calendar.monthrange(current_year, m)[1]
                dt = datetime(current_year, m, last_day)
                total = sum(i[1] * i[2] for i in items)
                order = create_synthetic_order(store_name, dt, total, items)
                if order: txns.append(order)
            
    except Exception as e:
        print(f"Error parsing Simula Excel {filepath}: {e}")
    return txns

def parse_frankie_excel(filepath, store_name="Frankie and Friends", current_year=2026):
    txns = []
    try:
        wb = openpyxl.load_workbook(filepath, data_only=True)
        for ws in wb.worksheets:
            current_order_no = None
            current_dt = None
            current_items = []
            
            for row in ws.iter_rows(min_row=2, values_only=True):
                if not row or not any(row): continue
                
                col0 = str(row[0]).strip()
                
                if col0.isdigit():
                    if row[1] and str(row[1]).isdigit() and row[2] and "202" in str(row[2]):
                        if current_items:
                            total = sum(i[1] * i[2] for i in current_items)
                            order = create_synthetic_order(store_name, current_dt, total, current_items, order_no=str(current_order_no))
                            if order: txns.append(order)
                            current_items = []
                            
                        current_order_no = str(row[1])
                        date_str = str(row[2]).strip()
                        current_dt = datetime(current_year, 1, 1)
                        for fmt in ("%d %b %Y %I:%M %p", "%Y-%m-%d %H:%M:%S"):
                            try:
                                current_dt = datetime.strptime(date_str, fmt)
                                break
                            except:
                                pass
                                
                    elif current_order_no:
                        prod = str(row[1]).strip()
                        price = clean_amount(row[2])
                        qty = clean_amount(row[4])
                        if qty == 0: qty = 1
                        current_items.append((prod, int(qty), price))
                        
            if current_items:
                total = sum(i[1] * i[2] for i in current_items)
                order = create_synthetic_order(store_name, current_dt, total, current_items, order_no=str(current_order_no))
                if order: txns.append(order)
                
    except Exception as e:
        print(f"Error parsing Frankie Excel {filepath}: {e}")
    return txns

def parse_pdf_report(filepath, store_name, current_year=2026):
    txns = []
    if not pdfplumber:
        return txns
        
    try:
        month = extract_month_from_filename(os.path.basename(filepath))
        if not month: month = 1
            
        last_day = calendar.monthrange(current_year, month)[1]
        dt = datetime(current_year, month, last_day)
        
        items = []
        
        with pdfplumber.open(filepath) as pdf:
            if store_name == "9 Matters":
                for page in pdf.pages:
                    tables = page.extract_tables()
                    for table in tables:
                        if not table: continue
                        for row in table[1:]:
                            if not row or len(row) < 5: continue
                            
                            row_str = str(row).upper()
                            if "TOTAL" in row_str or "NOTHING FOLLOWS" in row_str:
                                continue
                                
                            item_name = str(row[1]).strip()
                            if not item_name or item_name == "None": continue
                            
                            qty = clean_amount(row[2])
                            if qty == 0: qty = 1
                            amount = clean_amount(row[-1]) 
                            item_price = amount / qty if qty > 0 else amount
                            items.append((item_name, int(qty), item_price))
                            
            elif store_name == "Common Room":
                for page in pdf.pages:
                    text = page.extract_text()
                    if not text: continue
                    for line in text.split('\n'):
                        line = line.strip()
                        match = re.match(r'^\d+\.\s+(.*?)\s+(\d+)\s+([\d,]+\.\d{2})\s+[\d,]+\.\d{2}$', line)
                        if match:
                            prod_name = match.group(1).strip()
                            qty = clean_amount(match.group(2))
                            if qty == 0: qty = 1
                            amount = clean_amount(match.group(3))
                            item_price = amount / qty if qty > 0 else amount
                            items.append((prod_name, int(qty), item_price))

        if items:
            total_amount = sum(i[1] * i[2] for i in items)
            order = create_synthetic_order(store_name, dt, total_amount, items)
            if order: txns.append(order)
            
    except Exception as e:
        print(f"Error parsing PDF {filepath}: {e}")
    return txns

def parse_craft_central_html(filepath, current_year=2026):
    txns = []
    try:
        month = extract_month_from_filename(os.path.basename(filepath))
        
        items = []
        
        with open(filepath, 'r', encoding='utf-8') as f:
            soup = BeautifulSoup(f, 'html.parser')
            
            if not month:
                # Try to extract from text
                text = soup.get_text().upper()
                match = re.search(r'SALES FOR THE MONTH OF\s*([A-Z]+)', text)
                if match:
                    month_str = match.group(1)[:3].lower()
                    months_dict = {
                        "jan": 1, "feb": 2, "mar": 3, "apr": 4, "may": 5, "jun": 6,
                        "jul": 7, "aug": 8, "sep": 9, "oct": 10, "nov": 11, "dec": 12
                    }
                    month = months_dict.get(month_str)
                    
            if not month: month = 1
            
            tables = soup.find_all('table')
            
            for table in tables:
                rows = table.find_all('tr')
                found_in_table = False
                for row in rows:
                    cols = [td.get_text(strip=True) for td in row.find_all(['th', 'td'])]
                    if not cols or len(cols) < 4: continue
                    if "REA" in cols[0] and "REA" in cols[1]:
                        prod_name = cols[1]
                        qty = clean_amount(cols[2])
                        amount = clean_amount(cols[3])
                        if qty > 0 and amount > 0:
                            item_price = amount / qty
                            items.append((prod_name, int(qty), item_price))
                            found_in_table = True
                
                if found_in_table:
                    break # Stop after parsing the first table with actual data
                            
        if items:
            total_amount = sum(i[1] * i[2] for i in items)
            last_day = calendar.monthrange(current_year, month)[1]
            dt = datetime(current_year, month, last_day)
            order = create_synthetic_order("Craft Central", dt, total_amount, items)
            if order: txns.append(order)
            
    except Exception as e:
        print(f"Error parsing Craft Central HTML {filepath}: {e}")
    return txns
