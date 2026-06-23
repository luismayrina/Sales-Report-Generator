"""
Combined Sales Report Generator
Sources: Shopify orders/transactions CSVs + Shopee + Lazada Excel exports
Output:  Full_Sales_Report.xlsx matching Sample Reports.xlsx format
Sheets:  1) Daily Sales Summary  2) Real Time Inventory Report
         3) Offtake Report Summary  4) Offtake Report Summary per sku
"""

import csv, copy, os, re
import openpyxl
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime
from collections import defaultdict

# ── CONFIG ────────────────────────────────────────────────────────────────────
ORDERS_CSV       = "orders_export(1).csv"
TRANSACTIONS_CSV = "transactions_export.csv"
SHOPEE_XLSX      = "Shopee Transaction Report.xlsx"
LAZADA_XLSX      = "Lazada Transaction Report.xlsx"
SAMPLE_XLSX      = "Sample Reports.xlsx"
OUTPUT_XLSX      = "Full_Sales_Report.xlsx"

PAYMONGO_RATE    = 0.015   # 1.5% (confirmed from sample row #2425)
SHOPEE_FEE_RATE  = 0.0566  # ~5.66% commission (Shopee deducts from Grand Total)

# Month name → column index in Offtake sheets (col C=JAN=3, D=FEB=4, …)
MONTH_COL = {1:3,2:4,3:5,4:6,5:7,6:8,7:9,8:10,9:11,10:12,11:13,12:14}

BASE_DIR = os.path.dirname(os.path.abspath(__file__))

# ── HELPERS ───────────────────────────────────────────────────────────────────
def fval(s):
    try:
        return float(str(s).strip().replace(",","")) if s not in (None,"","–","-") else 0.0
    except:
        return 0.0

def parse_shopify_date(s):
    if not s or not str(s).strip():
        return None
    try:
        return datetime.strptime(str(s).strip()[:19], "%Y-%m-%d %H:%M:%S")
    except:
        return None

def parse_shopee_date(s):
    if not s or str(s).strip() in ("", "-"):
        return None
    try:
        return datetime.strptime(str(s).strip()[:16], "%Y-%m-%d %H:%M")
    except:
        return None

def parse_lazada_date(s):
    if not s or str(s).strip() == "":
        return None
    for fmt in ("%d %b %Y %H:%M", "%d %b %Y"):
        try:
            return datetime.strptime(str(s).strip(), fmt)
        except:
            pass
    return None

def copy_style(src, dst):
    if src.has_style:
        dst.font      = copy.copy(src.font)
        dst.fill      = copy.copy(src.fill)
        dst.border    = copy.copy(src.border)
        dst.alignment = copy.copy(src.alignment)
        dst.number_format = src.number_format

def bank_label(method, gateway="", status="paid"):
    if status != "paid":
        return "PENDING"
    combined = (method + " " + gateway).lower()
    if "paymongo" in combined:
        return "UBP"
    if "bank deposit" in combined:
        return "BANK DEPOSIT"
    if "manual" in combined:
        return "MANUAL"
    return method or ""

def normalize_sku(name):
    """Normalize product name to match SKU variant names."""
    n = str(name).lower()
    mapping = {
        "relaxing naturals": "Relaxing Naturals",
        "fresh bamboo": "Fresh Bamboo",
        "white tea": "White Tea",
        "cotton clean": "Cotton Clean",
        "quiet elegance": "Quiet Elegance",
        "charming home": "Charming Home",
        "bright mornings": "Bright Mornings",
        "mountain breeze": "Mountain Breeze",
        "beach escape": "Beach Escape",
        "candy cane": "Candy Cane",
        "apple cinnamon": "Apple Cinnamon",
        "easy drive": "Easy Drive",
        "creamy matcha": "Creamy Matcha",
        "white lavender": "White Lavender",
        "green tea": "Green Tea",
        "flushie": "Flushie 50ml",
        "lavender": "Lavender",
        "eucalyptus": "Eucalyptus",
        "peppermint": "Peppermint",
        "holiday hugs": "Holiday Hugs 200ml",
        "reed sticks": "Reed Sticks",
        "bamboo reed sticks": "Reed Sticks",
        "tropical paradise": "Tropical Paradise 200ml",
        "waterbased": "Waterbased",
        "clean air": "Clean Air",
        "fresh bamboo 500": "Fresh Bamboo 500ml",
        "white tea 500": "White Tea 500ml",
        "cotton clean 500": "Cotton Clean 500ml",
        "relaxing naturals 500": "Relaxing Naturals 500ml",
        "handwash": "Hand Wash Relaxing Naturals 500ml",
        "hand wash": "Hand Wash Relaxing Naturals 500ml",
        "reed diffuser refill": "White Lavender",  # mapped by variation
        "modern luxe": "Modern Luxe Reed Diffuser",
        "serene sanctuary": "Serene Sanctuary Reed Diffuser",
        "urban oasis": "Urban Oasis Reed Diffuser",
    }
    # Sort mapping keys by length (longest first) to prevent shorter strings from incorrectly capturing longer ones
    sorted_mapping = sorted(mapping.items(), key=lambda x: len(x[0]), reverse=True)
    for key, val in sorted_mapping:
        if key in n:
            return val
    return name.strip()


# ── LOAD SHOPIFY ──────────────────────────────────────────────────────────────
def load_shopify_transactions(path):
    txns = {}
    with open(path, newline="", encoding="utf-8-sig") as f:
        for row in csv.DictReader(f):
            name   = row["Name"].strip()
            status = row["Status"].strip().lower()
            entry  = txns.get(name)
            if entry is None or (status == "success" and entry["status"] != "success"):
                txns[name] = {
                    "gateway":  row["Gateway"].strip(),
                    "status":   status,
                    "amount":   fval(row["Amount"]),
                    "paid_at":  parse_shopify_date(row["Created At"]),
                }
    return txns

def load_shopify_orders(path, txns):
    rows, seen = [], set()
    with open(path, newline="", encoding="utf-8-sig") as f:
        for r in csv.DictReader(f):
            name = r["Name"].strip()
            if name in seen or not r["Financial Status"].strip():
                continue
            seen.add(name)
            txn     = txns.get(name, {})
            status  = r["Financial Status"].strip().lower()
            paid_at = parse_shopify_date(r["Paid at"]) or txn.get("paid_at")
            gateway = txn.get("gateway", r["Payment Method"].strip())
            is_pm   = "paymongo" in gateway.lower() or "paymongo" in r["Payment Method"].lower()
            total   = fval(r["Total"])
            disc    = fval(r["Discount Amount"])
            ship    = fval(r["Shipping"])
            subtotal= fval(r["Subtotal"])
            gross   = subtotal + disc   # before discount
            dep_dt  = paid_at if status == "paid" else None
            rows.append({
                "source":       "Shopify",
                "date":         parse_shopify_date(r["Created at"]),
                "customer":     r["Billing Name"].strip(),
                "order_no":     name,
                "gross":        gross,
                "shipping":     ship or None,
                "discount":     disc or None,
                "paymongo_fee": round(total * PAYMONGO_RATE, 2) if is_pm else None,
                "net":          total,
                "deposit_date": dep_dt,
                "bank":         bank_label(r["Payment Method"], gateway, status),
                "ref":          r["Payment Reference"].strip() or None,
                "status":       status,
                # for offtake
                "items":        [],   # filled later if needed
            })
    return rows

def load_shopify_items(orders_path):
    """Return {order_name: [(product_name, qty, price)]}"""
    items = defaultdict(list)
    with open(orders_path, newline="", encoding="utf-8-sig") as f:
        for r in csv.DictReader(f):
            items[r["Name"].strip()].append((
                r["Lineitem name"].strip(),
                int(fval(r["Lineitem quantity"])),
                fval(r["Lineitem price"]),
            ))
    return items


# ── LOAD SHOPEE ───────────────────────────────────────────────────────────────
def load_shopee(path):
    wb = load_workbook(path)
    ws = wb.active
    headers = [c.value for c in ws[1]]
    hi = {h: i for i, h in enumerate(headers)}

    orders, seen = [], set()
    all_rows = list(ws.iter_rows(min_row=2, values_only=True))

    # Group multi-product orders
    grouped = defaultdict(list)
    for row in all_rows:
        oid = row[hi["Order ID"]]
        grouped[oid].append(row)

    for oid, grp in grouped.items():
        first    = grp[0]
        status   = str(first[hi["Order Status"]]).strip()
        paid_time = parse_shopee_date(str(first[hi["Order Paid Time"]]))
        created   = parse_shopee_date(str(first[hi["Order Creation Date"]]))

        # Sum amounts across line items
        total_payment = sum(fval(r[hi["Total Buyer Payment"]]) for r in grp)
        total_fee     = sum(fval(r[hi["Service Fee"]]) for r in grp)
        total_disc    = sum(fval(r[hi["Total Discount(PHP)"]]) for r in grp)
        total_ship    = sum(fval(r[hi["Buyer Paid Shipping Fee"]]) for r in grp)

        # Grand Total from first row (it's shared per order)
        grand = fval(first[hi["Grand Total"]])

        is_cancelled = status.lower() == "cancelled"
        dep_dt = paid_time if not is_cancelled else None

        items = [(str(r[hi["Product Name"]]).strip(),
                  int(fval(r[hi["Quantity"]])),
                  fval(r[hi["Original Price"]])) for r in grp]

        orders.append({
            "source":       "Shopee",
            "date":         created,
            "customer":     str(first[hi["Receiver Name"]]).strip(),
            "order_no":     oid,
            "gross":        total_payment,
            "shipping":     total_ship or None,
            "discount":     total_disc or None,
            "paymongo_fee": None,
            "payment_fee":  total_fee or None,   # Shopee commission
            "net":          grand,
            "deposit_date": dep_dt,
            "bank":         "UBP" if not is_cancelled else "CANCELLED",
            "ref":          None,
            "status":       "cancelled" if is_cancelled else "paid",
            "items":        items,
        })
    return orders


# ── LOAD LAZADA ───────────────────────────────────────────────────────────────
def load_lazada(path):
    wb = load_workbook(path)
    ws = wb.active
    headers = [c.value for c in ws[1]]
    hi = {h: i for i, h in enumerate(headers)}

    grouped = defaultdict(list)
    for row in ws.iter_rows(min_row=2, values_only=True):
        onum = row[hi["orderNumber"]]
        grouped[onum].append(row)

    orders = []
    for onum, grp in grouped.items():
        first   = grp[0]
        status  = str(first[hi["status"]]).strip().lower()
        created = parse_lazada_date(str(first[hi["createTime"]]))
        delivered = parse_lazada_date(str(first[hi["deliveredDate"]]))

        total_paid = sum(fval(r[hi["paidPrice"]]) for r in grp)
        total_unit = sum(fval(r[hi["unitPrice"]]) for r in grp)
        total_sdisc = sum(fval(r[hi["sellerDiscountTotal"]]) for r in grp)
        total_ship = sum(fval(r[hi["shippingFee"]]) for r in grp)

        is_cancelled = status == "canceled"
        dep_dt = delivered if status in ("delivered", "confirmed") else None

        items = [(str(r[hi["itemName"]]).strip(),
                  1,
                  fval(r[hi["unitPrice"]])) for r in grp]

        orders.append({
            "source":       "Lazada",
            "date":         created,
            "customer":     str(first[hi["customerName"]]).strip(),
            "order_no":     str(onum),
            "gross":        total_unit,
            "shipping":     total_ship or None,
            "discount":     abs(total_sdisc) if total_sdisc < 0 else None,
            "paymongo_fee": None,
            "payment_fee":  None,
            "net":          total_paid,
            "deposit_date": dep_dt,
            "bank":         "LAZADA" if not is_cancelled else "CANCELLED",
            "ref":          None,
            "status":       "cancelled" if is_cancelled else ("paid" if dep_dt else "pending"),
            "items":        items,
        })
    return orders


# ── SHEET 1: DAILY SALES SUMMARY ─────────────────────────────────────────────
def write_daily_sales(ws_out, ws_sample, all_orders):
    # Copy header row
    for j, cell in enumerate(ws_sample[1], 1):
        dst = ws_out.cell(row=1, column=j, value=cell.value)
        copy_style(cell, dst)

    # Column widths
    for col_letter, dim in ws_sample.column_dimensions.items():
        ws_out.column_dimensions[col_letter].width = dim.width

    DATE_FMT  = "d-mmm"
    DATE2_FMT = "mm-dd-yy"
    MONEY_FMT = "#,##0.00"

    # Sort by date, then source
    source_order = {"Shopify": 0, "Shopee": 1, "Lazada": 2}
    sorted_orders = sorted(
        all_orders,
        key=lambda o: (o.get("date") or datetime(2000,1,1), source_order.get(o["source"], 9))
    )

    for ri, o in enumerate(sorted_orders, start=2):
        vals = [
            o["source"],           # A  au
            o.get("date"),         # B  Date
            o["customer"],         # C  Customer Name
            o["order_no"],         # D  Order No.
            o.get("gross"),        # E  Gross Amount
            o.get("shipping"),     # F  Plus Shipping fee
            o.get("payment_fee"),  # G  Less payment fee (Shopee commission)
            None,                  # H  Less shipping fee
            o.get("discount"),     # I  Less Discount
            None,None,None,None,None,None,None,None,  # J-Q
            o.get("paymongo_fee"), # R  Less Paymongo fee
            None,                  # S  EWT
            o.get("net"),          # T  Net Amount
            o.get("deposit_date"), # U  Payout/Date Deposit
            o.get("bank"),         # V  BANK
            None,                  # W  SI
            o.get("deposit_date"), # X  Date Deposit
            o.get("ref"),          # Y  Reference No.
            None,None,None,None,   # Z-AC
        ]
        for ci, v in enumerate(vals, 1):
            cell = ws_out.cell(row=ri, column=ci, value=v)
            if ci == 2 and v:
                cell.number_format = DATE_FMT
            elif ci == 21 and v:
                cell.number_format = DATE_FMT

# ── SHEET 2: REAL TIME INVENTORY ─────────────────────────────────────────────
def write_inventory(ws_out, ws_sample):
    """
    Copies sample structure and sets up headers for Finished Goods + Raw Materials.
    """
    for i, row in enumerate(ws_sample.iter_rows(), 1):
        for j, cell in enumerate(row, 1):
            dst = ws_out.cell(row=i, column=j, value=cell.value)
            copy_style(cell, dst)
    
    # Placeholder note
    ws_out.cell(row=3, column=1, value="[Pending Data Source for Finished Goods + Raw Materials]")


# ── SHEET 3: OFFTAKE REPORT SUMMARY ──────────────────────────────────────────
def write_offtake_summary(ws_out, ws_sample, all_orders):
    """
    Copies relevant rows and formulas from the sample and fills in actual totals.
    """
    # 1. Map where rows should go in the output
    # First row is header
    for j, cell in enumerate(ws_sample[1], 1):
        dst = ws_out.cell(row=1, column=j, value=cell.value)
        copy_style(cell, dst)
    
    out_row = 2
    channel_row_map = {}

    # 2. Iterate through sample and copy all rows
    for i, row in enumerate(ws_sample.iter_rows(min_row=2), 2):
        channel_name = str(row[0].value).strip() if row[0].value else ""
        
        for j, cell in enumerate(row, 1):
            dst = ws_out.cell(row=out_row, column=j, value=cell.value)
            copy_style(cell, dst)
            if cell.data_type == 'f':
                dst.value = cell.value 
        
        if channel_name:
            channel_row_map[channel_name] = out_row
        out_row += 1

    # 3. Compute and write totals
    channel_totals = defaultdict(lambda: defaultdict(float))
    channel_map = {"Shopify": "Shopify", "Shopee": "Shopee", "Lazada": "Lazada"}
    for o in all_orders:
        if o["status"] in ("cancelled",): continue
        src = channel_map.get(o["source"], o["source"])
        dt = o.get("date")
        if dt: channel_totals[src][dt.month] += o.get("net") or 0

    hdr = [c.value for c in ws_out[1]]
    month_actual_col = {}
    for ci, val in enumerate(hdr, 1):
        if isinstance(val, datetime): month_actual_col[val.month] = ci

    for channel, mdata in channel_totals.items():
        ridx = channel_row_map.get(channel)
        if not ridx: continue
        for month, total in mdata.items():
            cidx = month_actual_col.get(month)
            if cidx:
                ws_out.cell(row=ridx, column=cidx, value=round(total, 2))



# ── SHEET 4: OFFTAKE PER SKU ─────────────────────────────────────────────────
def write_offtake_per_sku(ws_out, ws_sample, all_orders):
    """
    Copies sample structure and fills in quantity sold per SKU per month
    across all channels (Shopify, Shopee, Lazada).
    """
    # Copy sample structure first
    for i, row in enumerate(ws_sample.iter_rows(values_only=True), 1):
        for j, val in enumerate(row, 1):
            ws_out.cell(row=i, column=j, value=val)

    # Find the header row (row 3: VARIANT, PRICE, JAN, FEB, ...)
    months = ["JAN","FEB","MAR","APR","MAY","JUN","JUL","AUG","SEP","OCT","NOV","DEC"]
    month_num = {m: i+1 for i, m in enumerate(months)}
    hdr_row = None
    for i, row in enumerate(ws_sample.iter_rows(values_only=True), 1):
        if row[0] == "VARIANT":
            hdr_row = i
            break
    if not hdr_row:
        return

    hdr = [c.value for c in ws_sample[hdr_row]]
    month_col = {}  # month_num -> col index (1-based)
    for ci, val in enumerate(hdr, 1):
        if val in months:
            month_col[month_num[val]] = ci

    # Build SKU row map: sku_name -> row_idx
    sku_row = {}
    for i, row in enumerate(ws_sample.iter_rows(min_row=hdr_row+1, values_only=True), hdr_row+1):
        if row[0]:
            sku_row[str(row[0]).strip()] = i

    # Count quantities from all orders
    sku_qty = defaultdict(lambda: defaultdict(int))  # sku -> month -> qty
    for o in all_orders:
        if o["status"] == "cancelled":
            continue
        dt = o.get("date")
        if not dt:
            continue
        month = dt.month
        for (item_name, qty, price) in o.get("items", []):
            sku = normalize_sku(item_name)
            # Try to find a match in sku_row
            matched = None
            for known_sku in sku_row:
                if known_sku.lower() == sku.lower():
                    matched = known_sku
                    break
            
            matched = matched or sku
            sku_qty[matched][month] += qty

    # Write quantities
    last_row_idx = max(sku_row.values()) if sku_row else hdr_row + 1

    # Write quantities
    for sku, month_data in sku_qty.items():
        row_idx = sku_row.get(sku)
        if not row_idx:
            # Dynamically create new row
            last_row_idx += 1
            row_idx = last_row_idx
            sku_row[sku] = row_idx
            ws_out.cell(row=row_idx, column=1, value=sku)
            # Copy styles
            if last_row_idx > hdr_row + 1:
                from sales_report_app.core.workbook_utils import copy_style
                for col in range(1, ws_out.max_column + 1):
                    src_cell = ws_out.cell(row=last_row_idx - 1, column=col)
                    dst_cell = ws_out.cell(row=row_idx, column=col)
                    copy_style(src_cell, dst_cell)
                    
        for month, qty in month_data.items():
            col_idx = month_col.get(month)
            if col_idx:
                existing = ws_out.cell(row=row_idx, column=col_idx).value
                ws_out.cell(row=row_idx, column=col_idx,
                            value=(existing or 0) + qty)


# ── MAIN ──────────────────────────────────────────────────────────────────────
def main():
    import glob
    import re
    from sales_report_app.core.physical_parser import load_physical_channels, load_py_data
    from sales_report_app.core.report_generator import generate_full_report

    print("📂  Loading Shopify transactions …")
    txn_path = os.path.join(BASE_DIR, TRANSACTIONS_CSV)
    if not os.path.exists(txn_path): txn_path = os.path.join(BASE_DIR, "docs", TRANSACTIONS_CSV)
    txns = load_shopify_transactions(txn_path)
    print(f"    {len(txns)} transaction records")

    print("📂  Loading Shopify orders …")
    orders_path = os.path.join(BASE_DIR, ORDERS_CSV)
    if not os.path.exists(orders_path): orders_path = os.path.join(BASE_DIR, "docs", ORDERS_CSV)
    shopify_orders = load_shopify_orders(orders_path, txns)
    shopify_items  = load_shopify_items(orders_path)
    for o in shopify_orders:
        o["items"] = shopify_items.get(o["order_no"], [])
    print(f"    {len(shopify_orders)} orders")

    print("📂  Loading Shopee transactions …")
    shopee_path = os.path.join(BASE_DIR, SHOPEE_XLSX)
    if not os.path.exists(shopee_path): shopee_path = os.path.join(BASE_DIR, "docs", SHOPEE_XLSX)
    shopee_orders = load_shopee(shopee_path)
    print(f"    {len(shopee_orders)} orders ({sum(1 for o in shopee_orders if o['status']!='cancelled')} active)")

    print("📂  Loading Lazada transactions …")
    lazada_path = os.path.join(BASE_DIR, LAZADA_XLSX)
    if not os.path.exists(lazada_path): lazada_path = os.path.join(BASE_DIR, "docs", LAZADA_XLSX)
    lazada_orders = load_lazada(lazada_path)
    print(f"    {len(lazada_orders)} orders")

    # Find the OFFTAKE REPORT dynamically
    offtake_xlsx = None
    search_paths = [os.path.join(BASE_DIR, "docs", "*OFFTAKE REPORT.xlsx"), os.path.join(BASE_DIR, "*OFFTAKE REPORT.xlsx")]
    for path_pattern in search_paths:
        matches = glob.glob(path_pattern)
        if matches:
            # Prefer the exact match if there are multiple, or just take the first
            offtake_xlsx = matches[0]
            break
            
    if not offtake_xlsx:
        print("⚠️ Warning: Could not find any file ending with 'OFFTAKE REPORT.xlsx'. PY data will be zero.")
        offtake_filename = "Unknown"
        current_year = 2026
        py_year = 2025
    else:
        # Extract the current year from the filename
        offtake_filename = os.path.basename(offtake_xlsx)
        match = re.search(r'(20\d{2})', offtake_filename)
        current_year = int(match.group(1)) if match else 2026
        py_year = current_year - 1

        print(f"📂  Loading Offtake Report ({offtake_filename})...")
        print(f"📅  Detected Reporting Year: {current_year} (PY: {py_year})")
        
    physical_orders = load_physical_channels(offtake_xlsx, current_year=current_year, external_docs_dir=os.path.join(BASE_DIR, "docs"))
    py_data = load_py_data(offtake_xlsx, py_year=py_year)

    all_orders = shopify_orders + shopee_orders + lazada_orders + physical_orders
    print(f"\n📊  Total combined orders: {len(all_orders)}")

    template_path = os.path.join(BASE_DIR, SAMPLE_XLSX)
    output_path = os.path.join(BASE_DIR, OUTPUT_XLSX)
    
    print("💾  Building output workbook …")
    success, result_path = generate_full_report(
        all_orders=all_orders,
        template_path=template_path,
        output_path=output_path,
        py_data=py_data,
        current_year=current_year,
        logger=print
    )

    if success:
        print(f"\n✅  Saved → {result_path}")
    else:
        print(f"\n❌  Failed to save report: {result_path}")

    # ── Summary Table ──
    print("\n" + "─"*72)
    print(f"{'Source':<15} {'Orders':>8} {'Paid Total':>14} {'Pending Total':>14}")
    print("─"*72)
    
    unique_sources = sorted(list(set(o["source"] for o in all_orders)))
    for src in unique_sources:
        grp = [o for o in all_orders if o["source"] == src]
        paid    = sum(o["net"] for o in grp if o["status"] not in ("cancelled","pending") and o.get("net"))
        pending = sum(o["net"] for o in grp if o["status"] == "pending" and o.get("net"))
        n = len(grp)
        print(f"{src:<15} {n:>8} {paid:>14,.2f} {pending:>14,.2f}")
    print("─"*72)
    total_paid    = sum(o["net"] for o in all_orders if o["status"] not in ("cancelled","pending") and o.get("net"))
    total_pending = sum(o["net"] for o in all_orders if o["status"] == "pending" and o.get("net"))
    print(f"{'TOTAL':<15} {len(all_orders):>8} {total_paid:>14,.2f} {total_pending:>14,.2f}")
    print("─"*72)


if __name__ == "__main__":
    main()
